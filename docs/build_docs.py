"""Сборка Word-документов из Markdown: единый стиль + pandoc. Запускать системным Python (нужен python-docx).

python docs/build_docs.py  ->  docs/*.docx  (потом docs/polish.ps1 доводит таблицы и оглавление в Word)
"""
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

D = Path(__file__).resolve().parent
SRC = D / "src"
REF = SRC / "ref.docx"

INK = RGBColor(0x16, 0x1A, 0x22)
TEAL = RGBColor(0x1F, 0x8A, 0x7A)
MUTED = RGBColor(0x5B, 0x64, 0x72)


def make_reference():
    doc = Document(SRC / "ref_default.docx")
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles
    spec = {
        "Normal": ("Calibri", 11, INK, False),
        "Body Text": ("Calibri", 11, INK, False),
        "First Paragraph": ("Calibri", 11, INK, False),
        "Compact": ("Calibri", 11, INK, False),
        "Title": ("Cambria", 26, INK, True),
        "Subtitle": ("Calibri", 13, MUTED, False),
        "Heading 1": ("Cambria", 18, TEAL, True),
        "Heading 2": ("Cambria", 14, INK, True),
        "Heading 3": ("Calibri", 12, INK, True),
        "Block Text": ("Calibri", 11, INK, False),
        "Image Caption": ("Calibri", 9.5, MUTED, False),
        "Table Caption": ("Calibri", 9.5, MUTED, False),
    }
    for name, (font, size, color, bold) in spec.items():
        try:
            s = st[name]
        except KeyError:
            continue
        s.font.name = font
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = bold
        rpr = s.element.get_or_add_rPr()
        rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if rfonts is not None:
            for k in ["ascii", "hAnsi", "cs", "eastAsia"]:
                rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + k, font)
            for k in ["asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"]:
                rfonts.attrib.pop("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + k, None)
        pf = s.paragraph_format
        if name.startswith("Heading"):
            pf.space_before = Pt(16 if name == "Heading 1" else 12)
            pf.space_after = Pt(6)
            pf.keep_with_next = True
        elif name in ("Normal", "Body Text", "First Paragraph"):
            pf.space_after = Pt(6)
            pf.line_spacing = 1.15
        elif name == "Title":
            pf.space_after = Pt(4)
        elif name == "Image Caption":
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(REF)


DOCS = [
    ("razbor.md", "Разбор_решений.docx", True),
    ("rech.md", "Текст_речи.docx", False),
    ("otchet.md", "Отчёт.docx", False),
    ("shpargalka.md", "Шпаргалка_к_защите.docx", True),
]


def contents_block(md: str) -> str:
    """Обычный список разделов вместо поля-оглавления: Word не нужно ничего обновлять при открытии."""
    body = md.split("---", 2)[2] if md.startswith("---") else md
    heads = [l[2:].strip() for l in body.splitlines() if l.startswith("# ")]
    if len(heads) < 3:
        return md
    nl = chr(10)
    block = nl + "**Содержание**" + nl + nl + nl.join("- " + re.sub(r"^(\d+)\.", lambda m: m.group(1) + chr(92) + ".", h) for h in heads) + nl + nl
    if md.startswith("---"):
        a, b, c = md.split("---", 2)
        return f"---{b}---{block}{c}"
    return block + md


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def style_tables(path: Path):
    """Таблицы: светлые границы, бирюзовая шапка, «зебра», шрифт 10, ширина на всю страницу."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(path)

    def shade(cell, hex_fill):
        tcpr = cell._tc.get_or_add_tcPr()
        for old in tcpr.findall(qn("w:shd")):
            tcpr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        tcpr.append(shd)

    for t in doc.tables:
        tblpr = t._tbl.tblPr
        for tag in ("w:tblBorders", "w:tblW", "w:tblCellMar"):
            for old in tblpr.findall(qn(tag)):
                tblpr.remove(old)
        w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct"); tblpr.append(w)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "D8DCE0")
            borders.append(b)
        tblpr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", "50"), ("bottom", "50"), ("left", "90"), ("right", "90")):
            m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), v); m.set(qn("w:type"), "dxa"); mar.append(m)
        tblpr.append(mar)
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                if i == 0:
                    shade(cell, "1F8A7A")
                elif i % 2 == 0:
                    shade(cell, "F3F5F8")
                for par in cell.paragraphs:
                    par.paragraph_format.space_after = Pt(0)
                    for r in par.runs:
                        r.font.size = Pt(10)
                        if i == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.save(path)


def main():
    make_reference()
    for src, out, toc in DOCS:
        md = (SRC / src).read_text(encoding="utf-8")
        tmp = SRC / ("_build_" + src)
        tmp.write_text(contents_block(md) if toc else md, encoding="utf-8")
        cmd = ["pandoc", str(tmp), "-o", str(D / out), "--reference-doc", str(REF),
               "--resource-path", str(SRC), "-f", "markdown+pipe_tables"]
        subprocess.run(cmd, check=True)
        tmp.unlink()
        style_tables(D / out)
        print("готово:", out)


if __name__ == "__main__":
    main()
