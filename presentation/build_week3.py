"""Презентация недели 3 (3–5 минут перед демонстрацией): схема системы и принятые решения.

python presentation/build_week3.py   ->  presentation/week3.pptx
Стиль — как у презентации недели 1 (build_deck.py). Все числа — из reports/week2/ (реальные запуски).
Текст для выступления — в заметках к слайдам.
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "week3.pptx"

DARK = RGBColor(0x14, 0x17, 0x1F)
DARK2 = RGBColor(0x1F, 0x24, 0x31)
INK = RGBColor(0x16, 0x1A, 0x22)
MUTED = RGBColor(0x5B, 0x64, 0x72)
CARD = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xD6, 0x45, 0x2F)
RED_L = RGBColor(0xF2, 0x8B, 0x7D)
RED_SOFT = RGBColor(0xFB, 0xE4, 0xE0)
AMBER = RGBColor(0xE9, 0xA8, 0x3B)
AMBER_SOFT = RGBColor(0xFC, 0xF0, 0xDA)
TEAL = RGBColor(0x1F, 0x8A, 0x7A)
TEAL_L = RGBColor(0x5F, 0xD0, 0xBC)
TEAL_SOFT = RGBColor(0xDD, 0xF0, 0xEC)
GREY = RGBColor(0xB8, 0xBF, 0xCA)
GREY_D = RGBColor(0x8F, 0x97, 0xA3)
ON_DARK = RGBColor(0xAE, 0xB6, 0xC4)
HEAD, BODY = "Cambria", "Calibri"
NB = " "

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, name="txt"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for t, o in (para if isinstance(para, list) else [(para, {})]):
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    tb.name = name
    return tb


def _fx3d(s, preset, fill):
    """Объём и тени как в презентации недели 1 (там их ставил fx.ps1 через PowerPoint; здесь — сразу в XML)."""
    from lxml import etree
    from pptx.oxml.ns import qn
    spPr = s._element.spPr
    for tag in ("a:effectLst", "a:scene3d", "a:sp3d"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    def shadow(alpha, blur, dist):
        ef = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(ef, qn("a:outerShdw"), blurRad=str(int(blur * 12700)), dist=str(int(dist * 12700)),
                              dir="5400000", algn="t", rotWithShape="0")
        c = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 1000)))
    dark = "{:02X}{:02X}{:02X}".format(*[int(v * 0.6) for v in (fill if fill is not None else RGBColor(0, 0, 0))])
    if preset == "tile":
        shadow(26, 16, 5)
    elif preset == "tile_dark":
        shadow(55, 18, 6)
    elif preset.startswith("card_"):
        shadow(55, 22, 10)
    elif preset == "sphere":
        shadow(35, 10, 4)
    sc = etree.SubElement(spPr, qn("a:scene3d"))
    if preset.startswith("card_"):
        rx, ry = {"card_a": (332, 12), "card_b": (336, 9), "card_c": (340, 6)}[preset]
        cam = etree.SubElement(sc, qn("a:camera"), prst="perspectiveFront", fov="2400000")
        etree.SubElement(cam, qn("a:rot"), lat=str(rx * 60000), lon=str(ry * 60000), rev="0")
    else:
        etree.SubElement(sc, qn("a:camera"), prst="orthographicFront")
    etree.SubElement(sc, qn("a:lightRig"), rig="threePt", dir="t")
    if preset == "sphere":
        w = int(s.width / 2)
        sp = etree.SubElement(spPr, qn("a:sp3d"), prstMaterial="plastic")
        etree.SubElement(sp, qn("a:bevelT"), w=str(w), h=str(w), prst="circle")
    elif preset.startswith("card_"):
        sp = etree.SubElement(spPr, qn("a:sp3d"), extrusionH=str(8 * 12700))
        etree.SubElement(sp, qn("a:bevelT"), w=str(5 * 12700), h=str(2 * 12700), prst="softRound")
        ec = etree.SubElement(sp, qn("a:extrusionClr"))
        etree.SubElement(ec, qn("a:srgbClr"), val=dark)
    else:
        sp = etree.SubElement(spPr, qn("a:sp3d"))
        inset = 5 if preset == "tile_dark" else 6
        etree.SubElement(sp, qn("a:bevelT"), w=str(inset * 12700), h=str(int(2.5 * 12700)), prst="softRound")


def shape(slide, kind, x, y, w, h, fill=CARD, line=None, lw=1.5, radius=None, name="shp", d3="auto"):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if d3 == "auto":
        d3 = ("sphere" if kind == MSO_SHAPE.OVAL else
              ("tile_dark" if fill in (DARK, DARK2) else "tile") if kind == MSO_SHAPE.ROUNDED_RECTANGLE else None)
    if d3:
        _fx3d(s, d3, fill)
    s.name = name
    return s


def label_in(s, lines, size, color=WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.word_wrap = True
    lines = lines if isinstance(lines, list) else [(lines, {})]
    for i, (tx, o) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = tx
        r.font.name = o.get("font", font)
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.color.rgb = o.get("color", color)
    return s


def title(slide, t, sub=None, dark=False):
    text(slide, 0.6, 0.4, 12.1, 0.7, t, size=28, font=HEAD, bold=True, color=WHITE if dark else INK, name="title")
    if sub:
        text(slide, 0.6, 1.1, 12.1, 0.45, sub, size=16, color=ON_DARK if dark else MUTED, name="sub")


LETTER = {"Назлы": "A", "Лия": "B", "Шырайлым": "C"}


def speaker(slide, who, dark=False):
    text(slide, 9.9, 7.02, 2.9, 0.3, f"говорит: {LETTER.get(who, who)}", size=11, color=GREY_D if dark else GREY,
         align=PP_ALIGN.RIGHT, name="who")


def source(slide, t, dark=False):
    text(slide, 0.6, 7.02, 9.0, 0.3, t, size=10, color=GREY_D if dark else GREY, name="src")


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def table(slide, x, y, w, rows, col_w, size=13, head_fill=INK, zebra=CARD, row_h=0.42, bold_rows=(), name="tbl"):
    gt = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    gt.name = name
    tb = gt.table
    for j, cw in enumerate(col_w):
        tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tb.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = head_fill if i == 0 else (zebra if i % 2 == 0 else WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if isinstance(val, tuple):
                val, col = val
            else:
                col = None
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = BODY
            r.font.size = Pt(size - (1 if i == 0 else 0))
            r.font.bold = i == 0 or i in bold_rows
            r.font.color.rgb = WHITE if i == 0 else (col or INK)
    return gt


def bar_chart(slide, x, y, w, h, cats, vals, colors, font_size=12, max_v=None, fmt="# ##0", name="chart"):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("знач.", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = BODY
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(font_size + 1)
    dl.font.bold = True
    dl.font.color.rgb = INK
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    if max_v:
        va.maximum_scale = max_v
    ca = ch.category_axis
    ca.reverse_order = True
    ca.tick_labels.font.size = Pt(font_size)
    ca.tick_labels.font.color.rgb = INK
    ca.format.line.fill.background()
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i]
    return gf




IMG = HERE / "img"


def chip(slide, x, y, w, h, t, fill=TEAL_SOFT, color=TEAL, size=12):
    c = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=0.3, d3=None)
    label_in(c, t, size, color=color, bold=True, font="Consolas")
    return c


def box(slide, x, y, w, h, head, body, fill=CARD, head_color=INK, size=12.5, line=None):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=0.06, line=line, lw=2)
    text(slide, x + 0.18, y + 0.12, w - 0.36, 0.4, head, size=size + 2.5, bold=True, color=head_color, font=HEAD)
    text(slide, x + 0.18, y + 0.55, w - 0.36, h - 0.6, body, size=size, color=MUTED)


def picture(slide, path, x, y, w):
    p = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    p.line.color.rgb = GREY
    p.line.width = Pt(1)
    return p


# =====================================================================
# 1. Титул
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
text(s, 0.8, 0.62, 7.3, 0.4, "ПРОЕКТ «АНТИФРОД» · НЕДЕЛЯ 3 · КОМАНДА ИЗ ТРЁХ УЧАСТНИЦ", size=13, color=ON_DARK, bold=True)
text(s, 0.8, 1.15, 7.2, 2.0, "Антифрод в реальном времени", size=40, font=HEAD, bold=True, color=WHITE)
text(s, 0.8, 2.75, 6.9, 1.6, "Модель второй недели стала сервисом: операция приходит по одной, решение — сразу. "
                             "Интерфейс оператора и панель мониторинга работают с моделью только через API.",
     size=17, color=ON_DARK)
stats = [("≈" + NB + "0,2" + NB + "с", "ответ на одну операцию (медиана; p95 ≈ 0,3 с)", TEAL_L),
         ("0", "расхождение с пакетным score.py на 300 операциях", WHITE),
         ("13 из 13", "мошеннических в первой серии из 200 операций holdout помечены, ложных — 0", WHITE)]
for i, (big, small, col) in enumerate(stats):
    x = 0.8 + i * 2.45
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.75, 2.25, 2.0, fill=DARK2, radius=0.08)
    text(s, x + 0.2, 4.87, 1.95, 0.7, big, size=28, font=HEAD, bold=True, color=col)
    text(s, x + 0.2, 5.6, 1.9, 1.1, small, size=11.5, color=ON_DARK)
for i, (l, role, quote, col, x, y, preset) in enumerate([
        ("A", "сервис", "«FastAPI и модель»", RED, 8.35, 1.0, "card_a"),
        ("B", "интерфейс", "«путь одной операции»", AMBER, 8.95, 2.3, "card_b"),
        ("C", "мониторинг", "«что видно на панели»", TEAL, 9.55, 3.6, "card_c")]):
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.35, 2.1, fill=col, radius=0.09, d3=preset)
    label_in(c, [(l + "  ·  " + role, {"size": 19}), (quote, {"size": 13, "font": BODY, "bold": False}),
                 ("**** **** **** 48" + str(31 + i), {"size": 11, "font": BODY, "bold": False})],
             20, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    c.text_frame.margin_left = Inches(0.3)
    c.text_frame.margin_top = Inches(0.2)
speaker(s, "Назлы", dark=True)
notes(s, "Здравствуйте. Это третья неделя нашего антифрода.\n\n"
         "Две недели наша модель работала с таблицей целиком: на вход файл, на выход файл. В банке операции приходят "
         "по одной, и решение нужно сразу. Поэтому на этой неделе мы сделали из модели небольшой сервис.\n\n"
         "Система из трёх частей: сервис на FastAPI, интерфейс оператора и панель мониторинга. Интерфейс и панель "
         "общаются с моделью только через API сервиса.\n\n"
         "Три числа. Одна операция обрабатывается примерно за 0,2 секунды. Сервис считает ровно так же, как наш "
         "пакетный score.py — расхождение ноль. И в первой серии из 200 операций, которых модель не видела, все 13 "
         "мошеннических помечены, ложных тревог нет.\n\n"
         "Я расскажу про устройство сервиса, Лия — про путь одной операции и интерфейс, Шырайлым — про панель и наши "
         "решения. Потом покажем всё вживую.")

# =====================================================================
# 2. Архитектура
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Схема системы: три части, связь только через API", "Каждая часть запускается своей командой в своём окне терминала")
box(s, 0.6, 1.7, 5.9, 1.25, "Интерфейс оператора · Streamlit :8501",
    "ввод операции, строка из holdout, серия N операций, сверка с фактом", fill=AMBER_SOFT, head_color=INK)
box(s, 6.8, 1.7, 5.9, 1.25, "Панель мониторинга · Streamlit :8502",
    "success rate, fraud rate, график, последние срабатывания; обновляется сама каждые 5 с", fill=TEAL_SOFT)
text(s, 0.6, 3.02, 12.1, 0.35, "▼   HTTP · JSON   ▼", size=13, color=GREY_D, bold=True, align=PP_ALIGN.CENTER, font="Consolas")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 3.4, 12.1, 1.45, fill=DARK, radius=0.06)
text(s, 0.85, 3.5, 7.0, 0.45, "Сервис скоринга · FastAPI :8000", size=18, bold=True, color=WHITE, font=HEAD)
text(s, 0.85, 3.98, 7.3, 0.8, "проверка входа (pydantic) → признаки по истории карты → модель v2 → порог → журнал",
     size=13, color=ON_DARK)
for i, ep in enumerate(["POST /predict", "POST /label", "GET /metrics", "GET /health"]):
    chip(s, 8.25 + (i % 2) * 2.2, 3.55 + (i // 2) * 0.6, 2.05, 0.45, ep, fill=DARK2, color=TEAL_L, size=12)
text(s, 0.6, 4.92, 12.1, 0.35, "▼   внутри сервиса   ▼", size=13, color=GREY_D, bold=True, align=PP_ALIGN.CENTER,
     font="Consolas")
box(s, 0.6, 5.3, 3.9, 1.25, "Модель v2 + код score.py", "src.pipeline.compute, LightGBM, калибровка; "
                                                        "загружается один раз при старте", size=12)
box(s, 4.7, 5.3, 3.9, 1.25, "История карт в памяти", "{cc_num: операции}: 1,72 млн операций до 02.12.2020 22:00 "
                                                     "+ всё, что пришло в сервис", size=12)
box(s, 8.8, 5.3, 3.9, 1.25, "Журнал · SQLite", "время, операция, оценка, решение, статус, код ошибки, метка",
    size=12)
shape(s, MSO_SHAPE.RECTANGLE, 0.6, 6.68, 12.1, 0.3, fill=RED_SOFT, d3=None)
text(s, 0.75, 6.7, 11.9, 0.28, "Интерфейс и панель не загружают модель и не читают журнал — только запросы к сервису",
     size=12, color=RED, bold=True)
source(s, "service/main.py · frontend/app.py · dashboard/app.py · README.md, раздел «Неделя 3»")
speaker(s, "Назлы")
notes(s, "Схема системы. Наверху две части, с которыми работает человек: интерфейс оператора и панель мониторинга. "
         "Обе на Streamlit.\n\n"
         "Они обращаются к сервису на FastAPI по HTTP — передают и получают JSON. У сервиса четыре маршрута: "
         "predict — оценить операцию, label — передать фактическую метку, metrics — данные для панели, health — "
         "проверить, что сервис жив.\n\n"
         "Внутри сервиса три вещи. Модель второй недели вместе с кодом признаков — тем же, что в score.py, "
         "загружается один раз при старте. История карт — словарь в памяти: для каждой карты её прошлые операции. "
         "И журнал в SQLite: каждый запрос с оценкой, решением, статусом и меткой.\n\n"
         "Главное правило: интерфейс и панель не трогают модель и журнал напрямую. Только через API.")

# =====================================================================
# 3. API
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Сервис на FastAPI: маршруты и ошибки", "Схемы запроса и ответа — pydantic; из них FastAPI сам строит страницу /docs")
rows = [["Маршрут", "Что принимает", "Что возвращает"],
        [("POST /predict", TEAL), "строку датасета без is_fraud", "transaction_id, fraud_score, is_fraud_pred, threshold"],
        [("POST /label", TEAL), "transaction_id, is_fraud", "подтверждение; 404 — нет такой операции"],
        [("GET /metrics", TEAL), "—", "сводка, ряды для графика, срабатывания, ошибки"],
        [("GET /health", TEAL), "—", "{\"status\": \"ok\"}, модель, порог"],
        [("PUT /threshold", GREY_D), "новый порог (доп.)", "старый и новый порог — без перезапуска"]]
table(s, 0.6, 1.7, 7.4, rows, [1.75, 2.45, 3.2], size=12, row_h=0.52)
text(s, 0.6, 5.0, 7.4, 0.35, "Ответ /predict (кроме обязательных полей — для оператора):", size=12.5, bold=True)
text(s, 0.6, 5.35, 7.4, 1.5,
     "{\"transaction_id\": \"085fa1a7-…\", \"fraud_score\": 1.0, \"is_fraud_pred\": 1,\n"
     " \"threshold\": 0.5, \"risk_level\": \"высокий\", \"expected_value\": 306.81,\n"
     " \"reasons\": [\"Сумма 304.81 у.е. — дороже 99% покупок в категории…\"],\n"
     " \"card_history\": 659, \"latency_ms\": 164.6}", size=11, font="Consolas", color=INK)
box(s, 8.3, 1.7, 4.4, 2.45, "422 — некорректный вход",
    "нет поля, текст вместо числа, неизвестная категория, сумма < 0. FastAPI отклоняет сам; мы пишем запрос в "
    "журнал как ошибку — success rate на панели падает.", fill=RED_SOFT, head_color=RED, size=12.5)
box(s, 8.3, 4.35, 4.4, 1.45, "500 — упала модель",
    "запрос в журнал со статусом «ошибка», ответ 500.", fill=CARD, size=12.5)
box(s, 8.3, 6.0, 4.4, 0.95, "/docs", "каждый маршрут можно вызвать вручную", fill=TEAL_SOFT, head_color=TEAL,
    size=12)
source(s, "service/main.py · tests/test_service.py (5 тестов: health, predict+label, 422 в журнале, 404, порог)")
speaker(s, "Назлы")
notes(s, "Маршруты сервиса. Predict принимает одну операцию в формате строки датасета, без колонки is_fraud, и "
         "возвращает номер операции, оценку, решение «фрод — не фрод» и порог. Для оператора мы добавили группу "
         "риска из второй недели, денежную пользу проверки, причины и время ответа.\n\n"
         "Label принимает фактическую метку для уже оценённой операции. Metrics отдаёт всё для панели. Health "
         "говорит, что сервис жив и модель загружена. Дополнительно можно поменять порог без перезапуска.\n\n"
         "Ошибки. Если не хватает поля или вместо суммы текст, FastAPI сам отвечает кодом 422 — это работа pydantic. "
         "Мы перехватываем такой запрос и пишем его в журнал как ошибку, поэтому success rate на панели падает. Если "
         "упадёт сама модель — код 500 и тоже запись в журнал.\n\n"
         "Страницу /docs FastAPI строит сам из схем — на демонстрации мы на ней покажем ошибку.")

# =====================================================================
# 4. Путь одной операции
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Путь одной операции: от кнопки до решения ≈ 0,2 с", "Тот же код признаков и модели, что в score.py — переписывать ничего не пришлось")
steps = [("1", "Кнопка", "интерфейс собирает JSON строки датасета без is_fraud", AMBER),
         ("2", "Проверка", "POST /predict; pydantic проверяет поля и типы, иначе 422", AMBER),
         ("3", "История карты", "из памяти — прошлые операции этой карты + новая", TEAL),
         ("4", "Признаки", "src.pipeline.compute: 35 признаков блоков base, a, b, c, e", TEAL),
         ("5", "Модель", "LightGBM → калибровка → вероятность p", TEAL),
         ("6", "Решение", "p ≥ 0,5 → «фрод»; запись в журнал; операция — в историю карты", RED),
         ("7", "Ответ и метка", "оценка, решение, причины; потом метка → POST /label", RED)]
for i, (n, h, d, col) in enumerate(steps):
    x = 0.6 + i * 1.75
    shape(s, MSO_SHAPE.OVAL, x + 0.5, 1.75, 0.6, 0.6, fill=col)
    text(s, x + 0.5, 1.85, 0.6, 0.45, n, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font=HEAD)
    if i < len(steps) - 1:
        shape(s, MSO_SHAPE.RIGHT_ARROW, x + 1.25, 1.93, 0.4, 0.24, fill=GREY, d3=None)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.55, 1.6, 2.0, fill=CARD, radius=0.08)
    text(s, x + 0.1, 2.65, 1.4, 0.4, h, size=13, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 3.1, 1.4, 1.4, d, size=11, color=MUTED, align=PP_ALIGN.CENTER)
box(s, 0.6, 4.85, 6.0, 2.0, "Почему FastAPI между интерфейсом и моделью",
    "модель и история карт загружаются один раз, а не в каждом окне; одна точка проверки входа и один журнал; "
    "интерфейс, панель и replay.py работают одинаково; интерфейс можно заменить, не трогая модель.",
    fill=TEAL_SOFT, head_color=TEAL, size=12.5)
box(s, 6.8, 4.85, 5.9, 2.0, "Проверка: сервис = score.py",
    "300 операций holdout по одной через сервис и пакетом через score_frame с той же историей — расхождение "
    "вероятностей 0, группы риска совпали на 100%. Команда: python -m service.check_batch.",
    fill=CARD, size=12.5)
source(s, "service/scoring.py · service/check_batch.py")
speaker(s, "Лия")
notes(s, "Пройду по шагам, что происходит, когда оператор нажимает кнопку.\n\n"
         "Первое: интерфейс собирает из полей формы JSON — строку датасета, без колонки is_fraud. Второе: отправляет "
         "её в predict. Pydantic проверяет, что все поля на месте и типы правильные, иначе сразу 422.\n\n"
         "Третье: у нас признаки по истории карты — например, сколько часов назад была крупная покупка. Поэтому "
         "сервис берёт из памяти прошлые операции этой карты и дописывает к ним новую. Четвёртое: считает признаки "
         "тем же кодом, что score.py. Пятое: модель и калибровка дают вероятность.\n\n"
         "Шестое: если вероятность не меньше 0,5 — решение «фрод». Запрос пишется в журнал, а операция становится "
         "историей своей карты. Седьмое: ответ возвращается в интерфейс, примерно за 0,2 секунды. Если операция из "
         "holdout, интерфейс после ответа отправляет фактическую метку в label.\n\n"
         "Мы проверили, что сервис считает так же, как пакетный score.py: 300 операций, расхождение ноль.")

# =====================================================================
# 5. Интерфейс оператора
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Интерфейс оператора", "Streamlit · http://localhost:8501 · вкладки: одна операция, серия, некорректный запрос")
picture(s, IMG / "operator.png", 0.6, 1.65, 6.9)
items = [("Ввод операции", "форма с полями датасета; значения можно поменять"),
         ("Строка из holdout", "по номеру, случайная, случайная честная или мошенническая"),
         ("Результат", "«ФРОД / НЕ ФРОД» цветом, оценка, порог, группа риска, причины"),
         ("Сверка с фактом", "после ответа — настоящая метка и верно ли решение; метка → /label"),
         ("Серия операций", "следующие N строк holdout по очереди, с паузой; счётчики пойманных и ложных"),
         ("Некорректный запрос", "текст вместо суммы, нет поля, неизвестная категория → 422")]
for i, (h, d) in enumerate(items):
    y = 1.7 + i * 0.85
    shape(s, MSO_SHAPE.OVAL, 7.8, y + 0.07, 0.28, 0.28, fill=AMBER)
    text(s, 8.25, y, 4.5, 0.35, h, size=14, bold=True, font=HEAD)
    text(s, 8.25, y + 0.35, 4.5, 0.5, d, size=11.5, color=MUTED)
source(s, "frontend/app.py · replay.py (та же серия из терминала)")
speaker(s, "Лия")
notes(s, "Интерфейс оператора. Операцию можно ввести вручную или взять строку из отложенной выборки — по номеру "
         "или случайную. Значения можно поменять перед отправкой.\n\n"
         "Результат выделен цветом: красный — фрод, зелёный — не фрод. Рядом оценка, порог, группа риска и причины — "
         "например, «сумма дороже 98% покупок в этой категории».\n\n"
         "Для строки из holdout после ответа показывается настоящая метка и верно ли решение. Метка не уходит в "
         "predict — только после ответа, отдельно в label.\n\n"
         "Вкладка «Серия» подаёт следующие N строк по очереди с паузой — так на панели видно, как меняются "
         "показатели. И отдельная вкладка, чтобы специально отправить некорректный запрос.")

# =====================================================================
# 6. Панель мониторинга
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Панель мониторинга", "Streamlit · http://localhost:8502 · данные только из GET /metrics, обновление каждые 5 с")
picture(s, IMG / "dashboard.png", 0.6, 1.65, 6.6)
rows = [["Показатель", "Как считаем", "Когда известен"],
        ["Success rate", "без ошибки / все запросы к /predict", "сразу"],
        ["Fraud rate, модель", "is_fraud_pred = 1 / успешные", "сразу"],
        ["Fraud rate, факт", "is_fraud = 1 / с меткой", "позже"]]
table(s, 7.45, 1.65, 5.3, rows, [1.6, 2.5, 1.2], size=11.5, row_h=0.5)
box(s, 7.45, 3.85, 5.3, 1.85, "Почему два fraud rate различаются",
    "ложные тревоги (FP) поднимают модельный, пропуски (FN) — фактический; метки приходят позже и не для всех. "
    "В первой серии они совпали: 13 из 13, ложных 0. При пороге 0,8 модель пропускает операцию с p = 0,70.",
    fill=AMBER_SOFT, size=11.5)
box(s, 7.45, 5.85, 5.3, 1.05, "Success rate ≠ качество модели",
    "это здоровье сервиса; качество — точность и чувствительность по меткам", fill=TEAL_SOFT, head_color=TEAL,
    size=11.5)
source(s, "dashboard/app.py · service/journal.py (compute_metrics)")
speaker(s, "Шырайлым")
notes(s, "Панель мониторинга. Она берёт данные только из маршрута metrics и обновляется сама каждые пять секунд.\n\n"
         "Наверху сводка: сколько запросов, success rate, fraud rate по модели и фактический. Ниже — точность, "
         "чувствительность, сколько денег фрода остановлено и пропущено, время ответа. Дальше графики во времени и "
         "таблица последних срабатываний с меткой.\n\n"
         "Определения. Success rate — доля запросов к predict, обработанных без ошибки. Это здоровье сервиса, а не "
         "качество модели: модель может ошибиться в ответе, но запрос всё равно успешный.\n\n"
         "Fraud rate по модели — доля операций, которые модель пометила как фрод. Он известен сразу, в момент "
         "операции. Фактический — доля настоящего фрода среди операций, для которых пришла метка. Он известен только "
         "потом, после проверки.\n\n"
         "Они различаются, потому что модель ошибается в обе стороны: ложные тревоги поднимают модельный, пропуски — "
         "фактический. В нашей первой серии они совпали — все 13 пойманы, ложных нет. Если поднять порог до 0,8, "
         "модель пропустит операцию с оценкой 0,7 — и они разойдутся.")

# =====================================================================
# 7. Решения и ограничения
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
title(s, "Принятые решения и известные ограничения", dark=True)
dec = [("Порог 0,5", "по калиброванной вероятности: «фрод вероятнее, чем нет». На valid точность 98,3%, "
                     "чувствительность 94,1%. Меняется без перезапуска"),
       ("История — словарь в памяти", "стартовая — 1,72 млн операций без меток; новые дописываются к своей карте"),
       ("Журнал — SQLite", "переживает перезапуск; панель считает по текущей сессии"),
       ("Streamlit", "команда пишет на Python; без JavaScript и CORS; автообновление — st.fragment"),
       ("Holdout от преподавателя", "своего неиспользованного периода нет; пока — 02.12–11.12.2020 из fraudTest, "
                                    "38 981 операция, 131 фрод")]
lim = [("≈ 0,2 с на операцию", "признаки пересчитываются по всей истории карты"),
       ("После перезапуска", "операции сессии в истории теряются"),
       ("Одна копия сервиса", "история не общая между процессами"),
       ("Лимит 0,3% в день", "в сервисе не применяется: решение по порогу, очередь аналитиков — отдельно")]
text(s, 0.7, 1.3, 6.4, 0.4, "Решения", size=18, bold=True, color=TEAL_L, font=HEAD)
for i, (h, d) in enumerate(dec):
    y = 1.8 + i * 0.85
    text(s, 0.7, y, 6.4, 0.9, [[(h + ". ", {"bold": True, "color": WHITE}), (d, {"color": ON_DARK})]], size=13)
text(s, 7.6, 1.3, 5.1, 0.4, "Ограничения", size=18, bold=True, color=AMBER, font=HEAD)
for i, (h, d) in enumerate(lim):
    y = 1.8 + i * 0.95
    text(s, 7.6, y, 5.1, 0.9, [[(h + ". ", {"bold": True, "color": WHITE}), (d, {"color": ON_DARK})]], size=13)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 6.15, 12.0, 0.75, fill=DARK2, radius=0.2)
text(s, 0.95, 6.27, 11.6, 0.5, [[("Демонстрация: ", {"bold": True, "color": TEAL_L}),
                                 ("запуск и /docs → честная операция → мошенническая → серия 200 → ошибка → итог",
                                  {"color": WHITE})]], size=15)
speaker(s, "Шырайлым", dark=True)
notes(s, "Наши решения. Порог 0,5 по откалиброванной вероятности: если модель считает фрод вероятнее, чем нет, — "
         "фрод. На проверочном периоде это точность 98% и чувствительность 94%. Порог можно поменять без перезапуска.\n\n"
         "История карт — словарь в памяти, как и предлагает задание. Журнал — SQLite. Интерфейс и панель на "
         "Streamlit: мы пишем на Python, и не нужен JavaScript.\n\n"
         "Отдельно про holdout. Периода, который мы не использовали ни в обучении, ни в сравнении версий, у нас не "
         "осталось. По заданию его выдаёт преподаватель — мы его подставим одной командой. Пока показываем на 9 сутках "
         "декабря 2020 года, на которых модель не обучалась.\n\n"
         "Ограничения: одна операция — 0,2 секунды, потому что признаки пересчитываются по всей истории карты. "
         "После перезапуска операции сессии теряются. И лимит 0,3% проверок в день из прошлых недель в сервисе не "
         "применяется — здесь решение по порогу.\n\n"
         "Теперь покажем всё вживую.")

prs.save(OUT)
print(f"OK -> {OUT}")
